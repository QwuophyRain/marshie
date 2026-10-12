"""Excel templates for students/teachers, and a forgiving importer for the data they bring back.

Sheet types the importer understands (it works out which is which from the header row):
  * Q&A      columns: Question, Answer  (+ Category, Keywords, Short Answer, Source, Gathered By, Class)
  * Facts    columns: Topic, Fact       (+ Category, Source, Gathered By, Class)
  * Data     any table the students collected (heights, temperatures, survey results...)
             The dataset name comes from the "About" sheet, or else the sheet/file name.
"""
from __future__ import annotations

import csv
import io
import re
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .safety import has_personal_contact, looks_inappropriate

PINK = "F9B7CF"
CREAM = "FFF6FA"
INK = "5B3A4B"
CATEGORIES = [
    "Arithmetic", "Mathematics", "English Language", "Literature", "Science", "Technology",
    "Ghana History", "World History", "Geography", "Arts & Creativity", "Music", "Health & Environment",
    "Computing & ICT", "Our School", "Fun Facts", "Other",
]
QA_COLS = ["Category", "Question", "Answer", "Keywords", "Short Answer (for quiz)", "Source", "Gathered By", "Class"]
FACT_COLS = ["Topic", "Fact", "Category", "Source", "Gathered By", "Class"]
MAX_ROWS = 5000
MAX_TEXT = 1500

_thin = Side(style="thin", color="E8C9D6")
_border = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)


# ============================================================================= templates
def _style_header(ws, headers: list[str], widths: list[int] | None = None) -> None:
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = Font(bold=True, color=INK, size=12)
        cell.fill = PatternFill("solid", fgColor=PINK)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _border
        ws.column_dimensions[get_column_letter(c)].width = (widths or [])[c - 1] if widths and c <= len(widths) else 24
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"


def _blank_area(ws, ncols: int, first: int, last: int, prefill: dict[int, str] | None = None) -> None:
    for r in range(first, last + 1):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=r, column=c, value=(prefill or {}).get(c))
            cell.border = _border
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if r % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=CREAM)


def _example(ws, row: int, values: list) -> None:
    for c, v in enumerate(values, 1):
        cell = ws.cell(row=row, column=c, value=v)
        cell.font = Font(italic=True, color="9A8791")
        cell.border = _border
        cell.alignment = Alignment(wrap_text=True, vertical="top")


def build_template(kind: str = "all", school: str = "", class_name: str = "", category: str = "",
                   columns: list[str] | None = None, dataset_name: str = "", rows: int = 50) -> bytes:
    """kind: 'qa' | 'facts' | 'data' | 'all'."""
    rows = max(10, min(int(rows or 50), 500))
    wb = Workbook()
    wb.remove(wb.active)
    want_qa, want_facts, want_data = kind in ("qa", "all"), kind in ("facts", "all"), kind in ("data", "all")

    # ---- Instructions
    ws = wb.create_sheet("Instructions")
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 110
    lines = [
        ("Marshie Learning Sheet" + (f" - {school}" if school else ""), True),
        (f"Class: {class_name}" if class_name else "", False),
        ("", False),
        ("1. Pick the sheet that fits what you found out (you don't have to fill every sheet).", False),
        ("2. Fill in the white/cream rows. The grey italic row is only an EXAMPLE - you can leave or delete it.", False),
        ("3. Write full, true sentences. Marshie will say your answer out loud, so keep it short and friendly.", False),
        ("4. Always write who gathered it and where you found it (book, teacher, experiment, survey).", False),
        ("5. Give your file to your teacher. Don't add phone numbers, addresses or private things about people.", False),
        ("", False),
        ("Sheets:", True),
        ("   Q&A    - a question and its answer.  Example: 'What is the capital of Ghana?' -> 'Accra.'", False),
        ("   Facts  - a topic and a fact about it. Example: 'Baobab tree' -> 'A baobab can live for over 1,000 years.'", False),
        ("   Data   - a table of numbers you measured or surveyed (plant heights, temperatures, favourite fruits...).", False),
        ("            Marshie can then answer 'What is the average height?' or 'Who is tallest?'.", False),
        ("", False),
        ("Keywords (optional) help Marshie find the answer: separate with commas, e.g.  capital, city, Accra", False),
        ("Short Answer (optional) is used when Marshie quizzes you: just the key word, e.g.  Accra", False),
    ]
    for r, (text, bold) in enumerate(lines, 1):
        cell = ws.cell(row=r, column=1, value=text)
        cell.font = Font(bold=bold, size=16 if r == 1 else 11, color=INK)
        cell.alignment = Alignment(wrap_text=True)

    # ---- hidden list of categories for the drop-down
    lists = wb.create_sheet("Lists")
    for i, cat in enumerate(CATEGORIES, 1):
        lists.cell(row=i, column=1, value=cat)
    lists.sheet_state = "hidden"
    cat_range = f"=Lists!$A$1:$A${len(CATEGORIES)}"
    pre_cat = category if category in CATEGORIES else ""

    def add_category_validation(sheet, col_letter: str, last_row: int) -> None:
        dv = DataValidation(type="list", formula1=cat_range, allow_blank=True, showErrorMessage=False)
        sheet.add_data_validation(dv)
        dv.add(f"{col_letter}3:{col_letter}{last_row}")

    if want_qa:
        qa = wb.create_sheet("Q&A")
        _style_header(qa, QA_COLS, [20, 42, 60, 24, 22, 24, 20, 14])
        _example(qa, 2, ["EXAMPLE: Science", "What do plants need to grow?", "Plants need sunlight, water, air and good soil to grow.",
                         "plants, grow, sunlight", "sunlight and water", "Science book, Class 4", "Ama", class_name or "Class 4"])
        _blank_area(qa, len(QA_COLS), 3, rows + 2, {1: pre_cat, 8: class_name})
        add_category_validation(qa, "A", rows + 2)

    if want_facts:
        fa = wb.create_sheet("Facts")
        _style_header(fa, FACT_COLS, [28, 70, 22, 26, 20, 14])
        _example(fa, 2, ["EXAMPLE: Baobab tree", "A baobab tree can live for more than a thousand years and store water in its trunk.",
                         "Science", "Visit to the school garden", "Kofi", class_name or "Class 4"])
        _blank_area(fa, len(FACT_COLS), 3, rows + 2, {3: pre_cat, 6: class_name})
        add_category_validation(fa, "C", rows + 2)

    if want_data:
        cols = [c.strip() for c in (columns or []) if c and c.strip()] or ["Name", "Item", "Value", "Date", "Notes"]
        da = wb.create_sheet("Data")
        _style_header(da, cols, [22] * len(cols))
        _example(da, 2, ["EXAMPLE"] + [""] * (len(cols) - 1))
        _blank_area(da, len(cols), 3, rows + 2)
        ab = wb.create_sheet("About")
        ab.column_dimensions["A"].width = 26
        ab.column_dimensions["B"].width = 60
        ab["A1"], ab["B1"] = "About this data", ""
        ab["A1"].font = Font(bold=True, size=14, color=INK)
        for r, (k, v) in enumerate([("Dataset name", dataset_name or "My class survey"),
                                    ("What is it about?", ""), ("Gathered by", ""), ("Class", class_name)], 3):
            ab.cell(row=r, column=1, value=k).font = Font(bold=True, color=INK)
            c = ab.cell(row=r, column=2, value=v)
            c.fill = PatternFill("solid", fgColor=CREAM)
            c.border = _border

    wb.move_sheet("Lists", offset=len(wb.sheetnames))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ============================================================================= importing
def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "yes" if v else "no"
    if isinstance(v, float):
        return str(int(v)) if v == int(v) else str(round(v, 6))
    if isinstance(v, (datetime, date)):
        return v.date().isoformat() if isinstance(v, datetime) else v.isoformat()
    return re.sub(r"\s+", " ", str(v)).strip()


def _key(h: str) -> str:
    return re.sub(r"[^a-z]", "", h.lower())


_ALIASES = {
    "question": "question", "questions": "question", "q": "question",
    "answer": "answer", "answers": "answer", "a": "answer", "response": "answer",
    "category": "category", "subject": "category", "topic": "topic", "fact": "fact", "facts": "fact",
    "keywords": "keywords", "keyword": "keywords", "tags": "keywords",
    "shortanswerforquiz": "short", "shortanswer": "short", "short": "short", "quizanswer": "short",
    "source": "source", "gatheredby": "added_by", "student": "added_by", "name": "added_by", "class": "class",
}


def _read_sheets(data: bytes, filename: str) -> list[tuple[str, list[list[str]]]]:
    if filename.lower().endswith((".csv", ".tsv")):
        text = data.decode("utf-8-sig", errors="replace")
        delim = "\t" if filename.lower().endswith(".tsv") else (";" if text.count(";") > text.count(",") else ",")
        return [(Path(filename).stem, [[_cell(c) for c in row] for row in csv.reader(io.StringIO(text), delimiter=delim)][:MAX_ROWS])]
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheets = []
    for ws in wb.worksheets[:20]:
        if ws.sheet_state != "visible":
            continue
        rows = []
        for row in ws.iter_rows(values_only=True):
            rows.append([_cell(c) for c in row])
            if len(rows) >= MAX_ROWS:
                break
        while rows and not any(rows[-1]):
            rows.pop()
        sheets.append((ws.title, rows))
    return sheets


def _find_header(rows: list[list[str]]) -> int | None:
    for i, row in enumerate(rows[:8]):
        keys = {_ALIASES.get(_key(c)) for c in row if c}
        if {"question", "answer"} <= keys or {"topic", "fact"} <= keys:
            return i
    return None


def parse_workbook(data: bytes, filename: str) -> dict:
    """Never raises for bad content: returns what could be read plus a list of problems for the teacher."""
    result = {"entries": [], "datasets": [], "sheets": [], "errors": [], "warnings": [], "skipped_examples": 0}
    try:
        sheets = _read_sheets(data, filename)
    except Exception as exc:
        result["errors"].append({"sheet": "", "row": 0, "message": f"I couldn't open that file ({exc}). Please upload an .xlsx or .csv file."})
        return result

    about_name = ""
    for title, rows in sheets:
        if title.strip().lower() == "about":
            for row in rows:
                if row and row[0].strip().lower() in ("dataset name", "name of data", "title") and len(row) > 1 and row[1]:
                    about_name = row[1]

    for title, rows in sheets:
        low = title.strip().lower()
        if low in ("instructions", "about", "lists", "readme") or not rows:
            continue
        h = _find_header(rows)
        if h is not None:
            cols = [_ALIASES.get(_key(c)) for c in rows[h]]
            is_facts = "topic" in cols and "fact" in cols and not ({"question", "answer"} <= set(cols))
            kind = "facts" if is_facts else "qa"
            added = 0
            for ri in range(h + 1, len(rows)):
                row, line = rows[ri], ri + 1
                if not any(row):
                    continue
                if any(c.upper().startswith("EXAMPLE") for c in row if c):
                    result["skipped_examples"] += 1
                    continue
                rec: dict[str, str] = {}
                for ci, name in enumerate(cols):
                    if name and ci < len(row) and row[ci]:
                        rec.setdefault(name, row[ci])
                if kind == "facts":
                    topic, fact = rec.get("topic", ""), rec.get("fact", "")
                    if not topic and not fact:
                        continue                      # blank template row (only Category/Class pre-filled)
                    if not topic or not fact:
                        result["errors"].append({"sheet": title, "row": line, "message": "Needs both a Topic and a Fact."})
                        continue
                    entry = {"question": [f"Tell me about {topic}", f"What is {topic}", f"Who is {topic}", topic],
                             "answer": fact, "keywords": topic, "category": rec.get("category") or "Our School"}
                else:
                    if not rec.get("question") and not rec.get("answer"):
                        continue                      # blank template row (only Category/Class pre-filled)
                    if not rec.get("question") or not rec.get("answer"):
                        result["errors"].append({"sheet": title, "row": line, "message": "Needs both a Question and an Answer."})
                        continue
                    entry = {"question": rec["question"], "answer": rec["answer"], "keywords": rec.get("keywords", ""),
                             "short": rec.get("short", ""), "category": rec.get("category") or "General"}
                entry["source"] = rec.get("source", "")
                entry["added_by"] = ", ".join(x for x in (rec.get("added_by"), rec.get("class")) if x)
                text = " ".join([entry["answer"], entry["keywords"] if isinstance(entry["keywords"], str) else "", rec.get("question", rec.get("topic", ""))])
                if looks_inappropriate(text):
                    result["errors"].append({"sheet": title, "row": line, "message": "Not added: please use kind, school-friendly words."})
                    continue
                if has_personal_contact(entry["answer"]):
                    result["errors"].append({"sheet": title, "row": line, "message": "Not added: it looks like a phone number or email. Please remove personal contact details."})
                    continue
                if len(entry["answer"]) > MAX_TEXT:
                    result["errors"].append({"sheet": title, "row": line, "message": f"Answer is too long (over {MAX_TEXT} characters). Please shorten it."})
                    continue
                if len(entry["answer"]) > 400:
                    result["warnings"].append({"sheet": title, "row": line, "message": "Long answer - Marshie will take a while to say it."})
                result["entries"].append(entry)
                added += 1
            result["sheets"].append({"name": title, "type": "Facts" if kind == "facts" else "Q&A", "rows": added})
            continue

        # ---- otherwise treat as a data table: first non-empty row = headers
        start = next((i for i, r in enumerate(rows) if sum(1 for c in r if c) >= 2), None)
        if start is None:
            result["warnings"].append({"sheet": title, "row": 0, "message": "Skipped: nothing I could recognise (no headers found)."})
            continue
        headers = [c or f"Column {i + 1}" for i, c in enumerate(rows[start])]
        while headers and headers[-1].startswith("Column "):
            headers.pop()
        body = []
        for ri in range(start + 1, len(rows)):
            row = rows[ri]
            if not any(row) or any(c.upper().startswith("EXAMPLE") for c in row if c):
                if any(row):
                    result["skipped_examples"] += 1
                continue
            body.append((row + [""] * len(headers))[: len(headers)])
        if not body:
            result["warnings"].append({"sheet": title, "row": 0, "message": "Skipped: the table has no data rows yet."})
            continue
        if any(looks_inappropriate(" ".join(r)) for r in body) or looks_inappropriate(" ".join(headers)):
            result["errors"].append({"sheet": title, "row": 0, "message": "Data sheet not added: it contains words that are not school-friendly."})
            continue
        name = about_name if (about_name and title.strip().lower() in ("data", "dataset", "sheet1")) else (
            title if title.strip().lower() not in ("data", "dataset", "sheet1") else (about_name or Path(filename).stem))
        result["datasets"].append({"name": name, "columns": headers, "rows": body})
        result["sheets"].append({"name": title, "type": f"Data table '{name}'", "rows": len(body)})
    if not result["entries"] and not result["datasets"] and not result["errors"]:
        result["warnings"].append({"sheet": "", "row": 0, "message": "I found nothing to learn. Did you fill in the template rows?"})
    return result


# ============================================================================= export
def export_workbook(entries: list[dict], datasets: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Q&A"
    _style_header(ws, QA_COLS, [20, 42, 60, 24, 22, 24, 20, 14])
    for r, e in enumerate(entries, 2):
        vals = [e["category"], e["question"], e["answer"], ", ".join(e.get("keywords", [])), e.get("short", ""),
                e.get("source", ""), e.get("added_by", ""), ""]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for ds in datasets:
        sheet = wb.create_sheet(re.sub(r"[\\/*?:\[\]]", "", ds["name"])[:28] or "Data")
        _style_header(sheet, ds["columns"])
        for r, row in enumerate(ds["rows"], 2):
            for c, v in enumerate(row, 1):
                sheet.cell(row=r, column=c, value=v)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
